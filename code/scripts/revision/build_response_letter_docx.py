"""Build the R2 response letter DOCX from the markdown draft."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def add_paragraph(document: Document, text: str) -> None:
    if text.startswith("- "):
        document.add_paragraph(text[2:].strip(), style="List Bullet")
    elif re.match(r"^\d+\.\s+", text):
        document.add_paragraph(re.sub(r"^\d+\.\s+", "", text).strip(), style="List Number")
    else:
        document.add_paragraph(text)


def add_markdown_table(document: Document, lines: list[str]) -> None:
    rows: list[list[str]] = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for row_idx, row in enumerate(rows):
        for col_idx in range(width):
            cell_text = row[col_idx] if col_idx < len(row) else ""
            paragraph = table.cell(row_idx, col_idx).paragraphs[0]
            run = paragraph.add_run(cell_text)
            if row_idx == 0:
                run.bold = True


def build_docx(markdown_path: Path, output_path: Path) -> None:
    text = markdown_path.read_text(encoding="utf-8")
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Times New Roman"
    styles["Title"].font.size = Pt(16)
    styles["Heading 1"].font.name = "Times New Roman"
    styles["Heading 1"].font.size = Pt(13)
    styles["Heading 2"].font.name = "Times New Roman"
    styles["Heading 2"].font.size = Pt(11.5)
    styles["Heading 3"].font.name = "Times New Roman"
    styles["Heading 3"].font.size = Pt(10.5)

    first_heading = True
    pending_lines: list[str] = []
    table_lines: list[str] = []

    def flush_pending() -> None:
        nonlocal pending_lines
        if pending_lines:
            add_paragraph(document, " ".join(line.strip() for line in pending_lines))
            pending_lines = []

    def flush_table() -> None:
        nonlocal table_lines
        if table_lines:
            add_markdown_table(document, table_lines)
            table_lines = []

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if line.startswith("|") and line.endswith("|"):
            flush_pending()
            table_lines.append(line)
            continue
        flush_table()
        if not line.strip():
            flush_pending()
            continue
        if line.startswith("# "):
            flush_pending()
            paragraph = document.add_paragraph(line[2:].strip(), style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            first_heading = False
        elif line.startswith("## "):
            flush_pending()
            document.add_paragraph(line[3:].strip(), style="Heading 1")
            first_heading = False
        elif line.startswith("### "):
            flush_pending()
            document.add_paragraph(line[4:].strip(), style="Heading 2")
            first_heading = False
        elif first_heading:
            flush_pending()
            document.add_paragraph(line.strip())
        elif line.startswith("- ") or re.match(r"^\d+\.\s+", line):
            flush_pending()
            add_paragraph(document, line.strip())
        else:
            pending_lines.append(line)

    flush_table()
    flush_pending()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    build_docx(args.input, args.output)


if __name__ == "__main__":
    main()
